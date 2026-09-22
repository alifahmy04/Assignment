from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def set_font(run, size=11, bold=False):
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def shade(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    p_pr.append(shading)


def add_heading(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.style = "Heading 1"
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    set_font(run, 13, True)
    return paragraph


def add_body(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.08
    run = paragraph.add_run(text)
    set_font(run)
    return paragraph


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
heading_style = styles["Heading 1"]
heading_style.font.name = "Aptos"
heading_style.font.color.rgb = RGBColor(0, 0, 0)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(4)
run = title.add_run("Elemental Anomaly Prototype Reflection")
set_font(run, 20, True)
run.font.color.rgb = RGBColor(0, 0, 0)

details = doc.add_paragraph()
details.alignment = WD_ALIGN_PARAGRAPH.CENTER
details.paragraph_format.space_after = Pt(14)
run = details.add_run("DMET 909 Assignment 1 | Ali Mohamed Fahmy Elsayed | ID 58-6345")
set_font(run, 10)

add_heading(doc, "Prototype Overview")
add_body(doc, "The Unity prototype represents the opening Fire Reactor area of Elemental Anomaly. It focuses on the first stage of the intended progression loop: the player begins with only a basic melee weapon, reaches a major Fire reactor, removes the suppression on Fire, absorbs energy from nearby environmental sources, and spends limited charges on a Fire attack. Restricting the prototype to one contained room kept the scope manageable while still demonstrating the main game concept.")

add_heading(doc, "Connection to the Game Concept")
add_body(doc, "The prototype supports the game concept by combining third-person exploration, combat, elemental progression, and resource management in one short playable sequence. The reactor acts as a clear progression milestone. Before interacting with it, the player can move, aim, and use melee attacks but cannot use Fire. After absorbing the reactor, Fire becomes available through a charge meter and a basic projectile attack. This mirrors the GDD idea that elemental powers are initially suppressed and are restored through powerful sources within the research facility.")

add_heading(doc, "Key Systems Demonstrated")
add_body(doc, "The scene includes a third-person camera, reticle-based aiming, player movement, melee combat, Fire projectiles, enemy health, player health, and environmental Fire vents. Fire vents can be absorbed after Fire is unlocked, temporarily lose their glow, and regenerate after 20 seconds. This gives the player a simple reason to consider when to spend Fire charges. The reactor also begins in a bright, energized state and becomes dark after its power is absorbed, making the progression change visible in the environment.")

add_heading(doc, "Visual and Level Design Choices")
add_body(doc, "The room uses an industrial research-facility layout with a central reactor, muted surfaces, imported science-fiction props, and warm Fire effects. The orange glow of the reactor and vents contrasts with the darker environment to guide attention toward interactive elemental objects. A small HUD communicates Fire charges, player health, the current combat mode, and interaction instructions. These choices support the intended atmosphere of escaping a controlled and hostile facility while gradually becoming more capable.")

add_heading(doc, "Reflection and Future Development")
add_body(doc, "The prototype shows that the core Fire loop is understandable in a small space: unlock an ability, collect its resource, and use it in combat. The main limitation is that the current enemy behaviour is deliberately simple, with contact damage rather than pursuit or attacks. If the project continued, I would add water and ice powers, elemental strengths and weaknesses, more varied enemies, additional facility sectors, sound effects, and a more complete objective flow. For this assignment, the Fire Reactor room provides a focused and playable proof of the larger Elemental Anomaly design.")

doc.core_properties.title = "Elemental Anomaly Prototype Reflection"
doc.core_properties.author = "Ali Mohamed Fahmy Elsayed"
doc.save("output/docx/Elemental_Anomaly_Prototype_Reflection.docx")
